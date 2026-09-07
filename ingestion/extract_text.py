from __future__ import annotations

import datetime as dt
import unicodedata
from dataclasses import asdict, dataclass
from pathlib import Path


class UnsupportedFileTypeError(ValueError):
    pass


class ExtractionError(Exception):
    pass


@dataclass
class ExtractedDocument:
    file_name: str
    file_path: str
    raw_text: str
    page_count: int | None
    extracted_at: str
    needs_ocr: bool = False

    def to_dict(self) -> dict:
        return asdict(self)


def extract_text(file_path: str) -> ExtractedDocument:
    """
    Input: đường dẫn tới file .pdf, .docx hoặc .txt
    Output: ExtractedDocument
    Raises: UnsupportedFileTypeError nếu không phải .pdf/.docx/.txt
            FileNotFoundError nếu file không tồn tại
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(str(path))
    suffix = path.suffix.lower()
    if suffix not in {".pdf", ".docx", ".txt"}:
        raise UnsupportedFileTypeError(f"Unsupported file type: {suffix}")
    if path.stat().st_size == 0:
        raise ExtractionError(f"Empty file: {path.name}")

    extracted_at = dt.datetime.now(dt.timezone.utc).isoformat()
    try:
        if suffix == ".pdf":
            return _extract_pdf(path, extracted_at)
        if suffix == ".txt":
            return _extract_txt(path, extracted_at)
        return _extract_docx(path, extracted_at)
    except (UnsupportedFileTypeError, FileNotFoundError, ExtractionError):
        raise
    except Exception as exc:
        raise ExtractionError(f"Cannot extract {path.name}: {exc}") from exc


def _extract_txt(path: Path, extracted_at: str) -> ExtractedDocument:
    raw_text = unicodedata.normalize("NFC", path.read_text(encoding="utf-8"))
    return ExtractedDocument(
        file_name=path.name,
        file_path=str(path.resolve()),
        raw_text=raw_text,
        page_count=None,
        extracted_at=extracted_at,
        needs_ocr=False,
    )


def _extract_pdf(path: Path, extracted_at: str) -> ExtractedDocument:
    import pdfplumber

    pages_text: list[str] = []
    page_count = 0
    with pdfplumber.open(path) as pdf:
        page_count = len(pdf.pages)
        if page_count == 0:
            raise ExtractionError(f"Corrupt or empty PDF: {path.name}")
        for page in pdf.pages:
            pages_text.append(page.extract_text() or "")
    raw_text = unicodedata.normalize("NFC", "\n".join(pages_text))
    needs_ocr = not raw_text.strip()
    return ExtractedDocument(
        file_name=path.name,
        file_path=str(path.resolve()),
        raw_text="" if needs_ocr else raw_text,
        page_count=page_count,
        extracted_at=extracted_at,
        needs_ocr=needs_ocr,
    )


def _extract_docx(path: Path, extracted_at: str) -> ExtractedDocument:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = Document(path)
    except Exception as exc:
        raise ExtractionError(f"Corrupt Word file: {path.name}: {exc}") from exc
    pieces: list[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = (Paragraph(child, doc).text or "").strip()
            if text:
                pieces.append(text)
        elif child.tag == qn("w:tbl"):
            text = _table_to_text(Table(child, doc))
            if text:
                pieces.append(text)
    raw_text = unicodedata.normalize("NFC", "\n".join(pieces))
    return ExtractedDocument(
        file_name=path.name,
        file_path=str(path.resolve()),
        raw_text=raw_text,
        page_count=None,
        extracted_at=extracted_at,
        needs_ocr=False,
    )


def _table_to_text(table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [" ".join((cell.text or "").split()) for cell in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)
