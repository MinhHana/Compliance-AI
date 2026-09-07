from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from ingestion.batch import process_directory
from ingestion.chunk_by_dieu import SourceMetadata, chunk_by_dieu
from ingestion.extract_text import ExtractedDocument, ExtractionError, extract_text

META = SourceMetadata(
    source_doc="Thông tư test",
    loai_van_ban="thong_tu",
    ngay_hieu_luc="2017-03-15",
    trang_thai="con_hieu_luc",
)

EMPTY_PDF = b"""%PDF-1.1
1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj
2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj
3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 3 3]>>endobj
xref
0 4
0000000000 65535 f 
0000000009 00000 n 
0000000052 00000 n 
0000000101 00000 n 
trailer<</Size 4/Root 1 0 R>>
startxref
178
%%EOF
"""


def _doc(text: str, **kw) -> ExtractedDocument:
    return ExtractedDocument(
        file_name=kw.get("file_name", "a.docx"),
        file_path="a.docx",
        raw_text=text,
        page_count=None,
        extracted_at="2026-01-01T00:00:00+00:00",
    )


def _write_docx(path: Path, text: str) -> None:
    doc = Document()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    doc.save(path)


def test_three_dieu():
    text = (
        "Điều 1. Phạm vi\nÁp dụng cho tổ chức tín dụng.\n\n"
        "Điều 2. Giải thích\nTCTD là tổ chức tín dụng.\n\n"
        "Điều 3. Nguyên tắc\nCho vay phải có bảo đảm."
    )
    chunks = chunk_by_dieu(_doc(text), META)
    assert len(chunks) == 3
    assert [c.dieu for c in chunks] == ["Điều 1", "Điều 2", "Điều 3"]


def test_scan_pdf_no_crash(tmp_path: Path):
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(EMPTY_PDF)
    extracted = extract_text(str(pdf))
    assert extracted.needs_ocr is True
    assert extracted.raw_text == ""


def test_free_text_one_chunk():
    chunks = chunk_by_dieu(_doc("Quy trình nội bộ không có cấu trúc điều."), META)
    assert len(chunks) == 1
    assert chunks[0].dieu is None


def test_batch_one_bad_file_continues(tmp_path: Path):
    raw = tmp_path / "raw"
    processed = tmp_path / "processed"
    raw.mkdir()
    good = raw / "ok.docx"
    _write_docx(
        good,
        "Điều 1. Một\nNội dung một.\nĐiều 2. Hai\nNội dung hai.\nĐiều 3. Ba\nNội dung ba.",
    )
    (good.with_suffix(good.suffix + ".meta.json")).write_text(
        '{"source_doc":"T","loai_van_ban":"thong_tu","ngay_hieu_luc":"2017-03-15","trang_thai":"con_hieu_luc","van_ban_thay_the":null}',
        encoding="utf-8",
    )
    bad = raw / "bad.pdf"
    bad.write_bytes(b"not a pdf")
    (bad.with_suffix(bad.suffix + ".meta.json")).write_text(
        '{"source_doc":"T","loai_van_ban":"thong_tu","ngay_hieu_luc":"2017-03-15","trang_thai":"con_hieu_luc","van_ban_thay_the":null}',
        encoding="utf-8",
    )
    results = process_directory(raw, processed)
    by_file = {r["file"]: r for r in results}
    assert by_file["ok.docx"]["ok"] is True
    assert by_file["ok.docx"]["chunks"] == 3
    assert by_file["bad.pdf"]["ok"] is False
    assert (processed / "ok.json").exists()


def test_long_dieu_splits_khoan():
    khoan = "\n".join(f"{i}. " + ("x" * 200) for i in range(1, 10))
    text = f"Điều 5. Cho vay\n{khoan}"
    assert len(text) > 1500
    chunks = chunk_by_dieu(_doc(text), META)
    assert len(chunks) > 1
    assert all(c.dieu == "Điều 5" for c in chunks)
    assert chunks[0].khoan is not None


def test_empty_file(tmp_path: Path):
    p = tmp_path / "empty.pdf"
    p.write_bytes(b"")
    with pytest.raises(ExtractionError):
        extract_text(str(p))


SOP_META = SourceMetadata(
    source_doc="Bảng phân chia thẩm quyền",
    loai_van_ban="sop_noi_bo",
    ngay_hieu_luc="2024-04-10",
    trang_thai="con_hieu_luc",
)


def test_docx_keeps_table_in_order(tmp_path: Path):
    path = tmp_path / "han_muc.docx"
    doc = Document()
    doc.add_paragraph("Bảng hạn mức")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Cấp"
    table.cell(0, 1).text = "Hạn mức"
    table.cell(1, 0).text = "Trưởng phòng giao dịch"
    table.cell(1, 1).text = "100 triệu VND"
    doc.add_paragraph("Hết bảng")
    doc.save(path)
    extracted = extract_text(str(path))
    text = extracted.raw_text
    assert "100 triệu VND" in text
    assert "Trưởng phòng giao dịch" in text
    assert text.index("Bảng hạn mức") < text.index("100 triệu")
    assert text.index("100 triệu") < text.index("Hết bảng")


def test_sop_splits_by_muc_and_phu_luc():
    text = (
        "BẢNG PHÂN CHIA THẨM QUYỀN (TRONG TÍN DỤNG)\n\n"
        "1. Định nghĩa\n"
        "Tổng hạn mức tín dụng là hạn mức tối đa được duyệt.\n\n"
        "2. Đối tượng áp dụng\n"
        "Pháp nhân và cá nhân.\n\n"
        "3. Nguyên tắc áp dụng\n"
        "Căn cứ giá trị tài sản đảm bảo và xếp hạng tín dụng.\n\n"
        "[PHỤ LỤC _ NGƯỜI CÓ THẨM QUYỀN THEO TỪNG BỘ PHẬN]\n"
        "Trưởng phòng giao dịch | 100 triệu VND\n"
        "Giám đốc chi nhánh | 500\n"
    )
    chunks = chunk_by_dieu(_doc(text), SOP_META)
    titles = [c.tieu_de_dieu or "" for c in chunks]
    assert len(chunks) >= 4
    assert any("Định nghĩa" in t for t in titles)
    assert any("Đối tượng" in t for t in titles)
    assert any("PHỤ LỤC" in t for t in titles)
    assert any("100 triệu VND" in c.noi_dung for c in chunks)
    assert all(c.dieu is None for c in chunks)


def test_table_row_not_section_heading():
    text = (
        "4. Giới hạn Quyền tự quyết\n"
        "Nội dung giới hạn.\n"
        "2. Thay đổi điều kiện |  | 500\n"
        "3. Gia hạn |  | 200\n"
        "5. Giảm Quyền Tự Quyết\n"
        "Không áp dụng với Trưởng phòng giao dịch.\n"
    )
    chunks = chunk_by_dieu(_doc(text), SOP_META)
    titles = [c.tieu_de_dieu or "" for c in chunks]
    assert any(t.startswith("4. Giới hạn") for t in titles)
    assert any(t.startswith("5. Giảm") for t in titles)
    assert not any("Thay đổi điều kiện" in t for t in titles)
    assert any("500" in c.noi_dung and "4. Giới hạn" in (c.tieu_de_dieu or "") for c in chunks)


def test_long_sop_section_packs_by_line():
    body = "\n".join(f"hàng bảng {i} | {i * 100}" for i in range(80))
    text = f"1. Hạn mức tín dụng riêng\n{body}"
    assert len(text) > 1500
    chunks = chunk_by_dieu(_doc(text), SOP_META)
    assert len(chunks) > 1
    assert all("Hạn mức tín dụng riêng" in (c.tieu_de_dieu or "") for c in chunks)
    assert "hàng bảng 0" in chunks[0].noi_dung
